# from fastapi import APIRouter, UploadFile, File, Form
# from app.services.ollama_client import get_key_points
# from app.services.ppt_generator import generate_ppt

# router = APIRouter()

# @router.post("/upload/")
# async def upload_transcript(
#     file: UploadFile = File(...),
#     prompt: str = Form("""
# You are a professional business analyst and corporate communications expert at Hitachi Digital Services.

# You are tasked with analyzing Microsoft Teams meeting transcripts. From the transcript, perform the following:

# 1. **Identify the meeting type** (e.g., Project Kickoff, Requirement Gathering, Technical Review, Client Demo, etc.).
# 2. **Summarize the meeting in 3–5 bullet points** with focus on agenda, key discussions, agreements, and tone.
# 3. **Extract all specific client requirements, decisions, action items, blockers, tools, and dependencies.**
# 4. **Decide the best PowerPoint format** to present the content to stakeholders (e.g., Timeline Deck, Executive Summary, Requirements Document, Proposal, etc.).
# 5. **Estimate the number of slides required.**
# 6. **For each slide:**
#    - Give a **meaningful title**.
#    - Provide **detailed and structured content**, suitable for slide use.
#    - The content should be multi-line, formatted clearly (bullets or short paragraphs).
#    - The information must come directly from the transcript — no placeholders like “Purpose, Date, Attendees.” Only useful, stakeholder-ready information.

# ---

# Output Format:
# ---

# **Meeting Type**:  
# **Summary**:  
# -  
# -  
# **Key Requirements / Action Points**:  
# -  
# -  
# **Recommended PPT Type**:  
# **Number of Slides**:  
# **Slide Breakdown**:  
# 1. **[Slide Title]**  
#    Content:  
#    - Bullet 1  
#    - Bullet 2  
#    - Bullet 3  
# 2. **[Slide Title]**  
#    Content:  
#    - Bullet 1  
#    - Bullet 2  
# ...
# n. **[Slide Title]**  
#    Content:  
#    ...
# Please ensure each slide’s content is detailed, concise, actionable, and derived directly from the meeting transcript. Avoid generic placeholder text. Write like you're preparing a deck for a real client meeting.
# """)):
#     # Read and decode uploaded transcript
#     content = (await file.read()).decode("utf-8")
    
#     # Extract key points using Ollama
#     key_points = get_key_points(content, prompt)
    
#     # Generate PowerPoint file
#     ppt_path = generate_ppt(key_points, filename=file.filename.split('.')[0])
    
#     return {
#         "message": "Presentation generated successfully.",
#         "key_points": key_points,
#         "ppt_file": ppt_path
#     }

# from fastapi import APIRouter, UploadFile, File, HTTPException
# from app.services.file_reader import read_file
# from app.services.groq_client import get_llm_response
# from app.services.ppt_generator import generate_ppt_from_content
# import os

# router = APIRouter()

# @router.post("/generate-summary")
# async def generate_summary(file: UploadFile = File(...)):
#     return await process_file(file, task="summary")

# @router.post("/generate-notes")
# async def generate_notes(file: UploadFile = File(...)):
#     return await process_file(file, task="notes")

# @router.post("/generate-ppt")
# async def generate_ppt_content(file: UploadFile = File(...)):
#     return await process_file(file, task="ppt")


# # Common logic reused by all endpoints
# async def process_file(file: UploadFile, task: str):
#     try:
#         content = await file.read()
#         raw_text = read_file(file.filename, content)

#         if not raw_text:
#             raise HTTPException(status_code=400, detail="File is empty or unsupported format.")

#         if task == "summary":
#             prompt = "Generate a concise summary (250 - 350 words) of the following content, capturing the main ideas and key points."
#         elif task == "notes":
#             prompt = "Create detailed bullet point notes from the following content, organizing key information clearly and concisely."
#         elif task == "ppt":
#             prompt = """
# You are a professional business analyst and corporate communications expert at Hitachi Digital Services.

# You are tasked with analyzing Microsoft Teams meeting transcripts. From the transcript, perform the following:

# 1. **Identify the meeting type** (e.g., Project Kickoff, Requirement Gathering, Technical Review, Client Demo, etc.).
# 2. **Summarize the meeting in 3–5 bullet points** with focus on agenda, key discussions, agreements, and tone.
# 3. **Extract all specific client requirements, decisions, action items, blockers, tools, and dependencies.**
# 4. **Decide the best PowerPoint format** to present the content to stakeholders (e.g., Timeline Deck, Executive Summary, Requirements Document, Proposal, etc.).
# 5. **Estimate the number of slides required.**
# 6. **For each slide:**
#    - Give a **meaningful title** Dont repeat the title names.
#    - Provide **detailed and structured content**, suitable for slide use.
#    - The content should be multi-line, formatted clearly (bullets or short paragraphs).
#    - The information must come directly from the transcript — no placeholders like “Purpose, Date, Attendees.” Only useful, stakeholder-ready information.

# ---

# Output Format:
# ---

# **Meeting Type**:  
# **Summary**:  
# -  
# -  
# **Key Requirements / Action Points**:  
# -  
# -  
# **Recommended PPT Type**:  
# **Number of Slides**:  
# **Slide Breakdown**:  
# 1. **[Slide Title]**  
#    Content:  
#    - Bullet 1  
#    - Bullet 2  
#    - Bullet 3  
# 2. **[Slide Title]**  
#    Content:  
#    - Bullet 1  
#    - Bullet 2  
# ...
# n. **[Slide Title]**  
#    Content:  
#    ...
# Please ensure each slide’s content is detailed, concise, actionable, and derived directly from the meeting transcript. Avoid generic placeholder text. Write like you're preparing a deck for a real client meeting.
# """
#         else:
#             raise HTTPException(status_code=400, detail="Invalid task type.")

#         result = get_llm_response(prompt, raw_text)

#         if task == "ppt":
#             ppt_path = generate_ppt(result.splitlines(), filename=os.path.splitext(file.filename)[0])
#             return {"message": "PPT generated successfully.", "ppt_file": ppt_path}

#         return {"message": f"{task.capitalize()} generated successfully.", "content": result}

#     except Exception as e:
#         raise HTTPException(status_code=500, detail=str(e))


from fastapi import APIRouter, UploadFile, File, HTTPException
from fastapi.responses import FileResponse
import os
import docx
from app.services.groq_client import get_llm_response, ppt_llm
from app.services.ppt_generator import create_presentation
from app.state import stored_text

router = APIRouter()

@router.post("/upload/")
async def upload_file(file: UploadFile = File(...)):
    global stored_text
    ext = os.path.splitext(file.filename)[1].lower()
    content = ""

    if ext == ".txt":
        content = (await file.read()).decode("utf-8")
    elif ext == ".docx":
        doc = docx.Document(await file.read())
        content = "\n".join([para.text for para in doc.paragraphs])
    else:
        raise HTTPException(status_code=400, detail="Only .txt or .docx files allowed.")

    stored_text = content.strip()
    if not stored_text:
        raise HTTPException(status_code=400, detail="File content is empty.")

    return {"message": "File uploaded and text extracted successfully."}


@router.get("/generate_summary/")
def generate_summary():
    if not stored_text:
        raise HTTPException(status_code=400, detail="No text uploaded yet.")
    prompt = """
            You are a professional meeting summarizer.
            
            Your task is to read the following transcript and write a concise summary of the meeting.
            Include the key discussion points, goals, decisions, and outcomes.
            Ensure the summary is clear, neutral, and limited to 150–200 words.
            
            Do not list individual speakers or timestamps.

            """
    return {"summary": get_llm_response(prompt, stored_text)}


@router.get("/generate_notes/")
def generate_notes():
    if not stored_text:
        raise HTTPException(status_code=400, detail="No text uploaded yet.")
    prompt = """
            You are a business analyst creating bullet point notes from a meeting transcript.
            
            Read the transcript and extract key discussion points, decisions, goals, technical ideas, and action items.
            Format them as structured, detailed bullet points.
            Avoid including speaker names or timestamps.
            Group points logically if multiple themes are discussed.
            
            Here is the transcript:
            """
    return {"notes": get_llm_response(prompt, stored_text)}


@router.get("/generate_ppt/")
def generate_ppt_api():
    if not stored_text:
        raise HTTPException(status_code=400, detail="No text uploaded yet.")

    prompt = """
        You are a professional business analyst and corporate communications expert at Hitachi Digital Services.
        
        You are tasked with analyzing a Microsoft Teams meeting transcript and generating a comprehensive PowerPoint presentation that includes every detail from the meeting. The PPT must cover the entire meeting content, including all key requirements, decisions, action items, blockers, tools, dependencies, and any other mentions, without omitting any important information. Do not generate a summary or key points separately—embed all details directly into the PPT slides.
        
        From the transcript, perform the following:
        1. **Identify the meeting type** (e.g., Project Kickoff, Requirement Gathering, Technical Review, Client Demo, etc.) and include it in a slide.
        4. **Decide the best PowerPoint format** to present the content to stakeholders (e.g., Timeline Deck, Executive Summary, Requirements Document, Proposal, etc.).
        5. **Estimate the number of slides required, ensuring each slide covers distinct aspects of the meeting, with no fixed limit on content per slide.**
        6. **For each slide:**
        - Give a **meaningful title** reflecting the slide's content (e.g., 1. **Meeting Type Identification**).
        - Provide **detailed and structured content**, suitable for slide use, capturing every detail from the transcript without restriction on the number of points.
        - The content should be multi-line, formatted clearly (bullets or short paragraphs), with the number of points varying based on the transcript's details (no minimum or maximum limit, but at least one point per slide).
        - Avoid including specific timestamps (e.g., 'at 17:17') unless they are critical to understanding the content; focus on the information itself.
        - The information must come directly from the transcript—no placeholders like “Purpose, Date, Attendees.” Only include stakeholder-ready, meeting-specific information.
        - Cover the entire meeting, including all key requirements, decisions, action items, blockers, tools, dependencies, and any other mentions, distributing content across slides to ensure comprehensive coverage.
        
        ---
        Output Format:
        ---
        **Recommended PPT Type**:  
        **Number of Slides**:  
        **Slide Breakdown**:  
        1. **Slide Title**  
        Content:  
        - Detail 1  
        - Detail 2  
        ...
        ...
        2. **Slide Title**  
        Content:  
        - Detail 1  
        - Detail 2  
        ...
        ...  
        ...  
        n. **Slide Title**  
        Content:  
        - Detail 1  
        - Detail 2  
        ...
        ...
        
        Here is the meeting transcript:\n\n
        """
 
    
    prompt += stored_text
    ai_output = ppt_llm(prompt)
    
    template_path = "custom--template (1).pptx"
    output_path = "static/presentations/output_presentation.pptx"
    presentation_title = "Hitachi Digital Services"
    presentation_subtitle = "Meeting Presentation"

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    create_presentation(ai_output, template_path, output_path, presentation_title, presentation_subtitle)

    return {"message": "PPT generated successfully.", "ppt_file": os.path.join("static", "presentations", "output_presentation.pptx")}


@router.get("/static/presentations/{file_path:path}")
async def serve_ppt(file_path: str):
    file_location = os.path.join("static", "presentations", file_path)
    if os.path.exists(file_location):
        return FileResponse(
            file_location,
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            filename=file_path,
            headers={"Content-Disposition": "attachment; filename=presentation.pptx"}
        )
    raise HTTPException(status_code=404, detail="File not found")
